"""`Triton.exe --selftest <folder>`: converts generated sample files through the
real engines and writes PASS/FAIL lines to <folder>/selftest.txt.

The exe is windowed (no stdout), and a PyInstaller bundle can miss data files
that only matter at runtime, so this is how a build is checked end to end.
"""

from __future__ import annotations

import datetime as dt
import sys
import traceback
from pathlib import Path

TR = "Çalışkan öğrenciler İstanbul'da şişman ğ ı"


def _samples(d: Path) -> dict[str, Path]:
    from docx import Document
    import openpyxl
    from odf.opendocument import OpenDocumentText
    from odf.text import H, P
    from PIL import Image
    from pptx import Presentation

    png = d / "resim.png"
    Image.new("RGB", (120, 80), (31, 143, 107)).save(png)

    doc = Document()
    doc.add_heading("Rapor", 1)
    doc.add_paragraph(TR)
    doc.save(str(d / "rapor.docx"))

    wb = openpyxl.Workbook()
    wb.active.append(["Ad", "Tarih"])
    wb.active.append(["Şebnem", dt.date(2026, 1, 2)])
    wb.save(d / "tablo.xlsx")

    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "Sunum " + TR
    prs.save(str(d / "sunum.pptx"))

    odt = OpenDocumentText()
    odt.text.addElement(H(outlinelevel=1, text="Özet"))
    odt.text.addElement(P(text=TR))
    odt.save(str(d / "ozet.odt"))

    (d / "not.txt").write_bytes(TR.encode("cp1254"))
    (d / "veri.csv").write_bytes("Ad;Tutar\nŞebnem;12,5\n".encode("cp1254"))
    return {p.name: p for p in d.iterdir() if p.is_file()}


def run(folder: str) -> int:
    from app.core.converter_registry import convert_file
    import pymupdf

    work = Path(folder)
    work.mkdir(parents=True, exist_ok=True)
    report = work / "selftest.txt"
    lines: list[str] = []
    failed = 0

    def record(name: str, ok: bool, detail: str = "") -> None:
        nonlocal failed
        failed += 0 if ok else 1
        lines.append(f"{'PASS' if ok else 'FAIL'} {name} {detail}".rstrip())
        report.write_text("\n".join(lines), encoding="utf-8")

    try:
        files = _samples(work)
    except Exception:  # noqa: BLE001
        record("build samples", False, traceback.format_exc())
        return 1

    def pdf_text(p: Path) -> str:
        with pymupdf.open(p) as d:
            return "".join(page.get_text() for page in d)

    jobs = [
        ("rapor.docx", "pdf", "Çalışkan"), ("tablo.xlsx", "pdf", "Şebnem"), ("sunum.pptx", "pdf", "Sunum"),
        ("ozet.odt", "pdf", "Çalışkan"), ("not.txt", "pdf", "şişman"), ("veri.csv", "pdf", "Şebnem"),
        ("resim.png", "jpg", None), ("resim.png", "pdf", None),
    ]
    for src_name, ext, needle in jobs:
        out = work / f"{Path(src_name).stem}_{Path(src_name).suffix[1:]}.{ext}"
        try:
            convert_file(files[src_name], out)
            ok = out.exists() and out.stat().st_size > 0
            if ok and ext == "pdf" and needle:
                ok = needle in pdf_text(out)
            record(f"{src_name} -> {ext}", ok)
        except Exception:  # noqa: BLE001
            record(f"{src_name} -> {ext}", False, traceback.format_exc().splitlines()[-1])

    # engines that live in other bundled libraries (opencv/numpy via pdf2docx, PyMuPDF rendering)
    try:
        pdf = work / "rapor_docx.pdf"
        convert_file(pdf, work / "rapor_pdf.docx")
        record("pdf -> docx (pdf2docx)", (work / "rapor_pdf.docx").stat().st_size > 0)
        convert_file(pdf, work / "rapor_pdf.png")
        record("pdf -> png (PyMuPDF)", (work / "rapor_pdf.png").stat().st_size > 0)
    except Exception:  # noqa: BLE001
        record("pdf -> docx / png", False, traceback.format_exc().splitlines()[-1])

    lines.append(f"DONE failed={failed}")
    report.write_text("\n".join(lines), encoding="utf-8")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(run(sys.argv[1]))
