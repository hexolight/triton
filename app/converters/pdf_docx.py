"""PDF -> DOCX built on PyMuPDF (reading) and python-docx (writing).

Keeps text with its font, size, bold/italic and colour, embedded images and detected tables, in reading
order, one Word page per PDF page. Layout is flowed rather than pixel-placed: columns become a single
column, text boxes become paragraphs. That is what keeps the app small -- the usual alternative
(pdf2docx) drags OpenCV and NumPy into the bundle, roughly 170 MB for one conversion direction.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Emu, Pt, RGBColor

_SUBSET = re.compile(r"^[A-Z]{6}\+")
_STYLE_WORDS = re.compile(r"[-,]?(Bold|Italic|Oblique|Regular|Roman|Light|Medium|Semibold|Black|MT|PS)+$", re.I)
_DOCX_IMAGES = {"png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff"}
_MARGIN = 36.0  # points; the page is sized to the PDF's own page, so a slim margin keeps text on one page


def _family(font: str) -> str | None:
    name = _SUBSET.sub("", font or "")
    name = _STYLE_WORDS.sub("", name).replace("-", " ").strip()
    return name or None


def _inside(box: pymupdf.Rect, bbox: tuple[float, float, float, float]) -> bool:
    cx, cy = (bbox[0] + bbox[2]) / 2, (bbox[1] + bbox[3]) / 2
    return box.x0 <= cx <= box.x1 and box.y0 <= cy <= box.y1


def _gap(y0: float, prev_bottom: float | None) -> Pt:
    return Pt(min(max((y0 - prev_bottom) if prev_bottom is not None else 0, 0), 24))


def _add_text(doc, block: dict, page_w: float, prev_bottom: float | None, left_edge: float) -> float:
    para = doc.add_paragraph()
    fmt = para.paragraph_format
    x0, y0, x1, y1 = block["bbox"]
    fmt.space_after = Pt(0)
    fmt.space_before = _gap(y0, prev_bottom)

    if abs((x0 + x1) / 2 - page_w / 2) < 6 and x0 - left_edge > 24:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif x0 - left_edge > 12:
        fmt.left_indent = Pt(min(x0 - left_edge, page_w / 2))

    lines = block.get("lines", [])
    for li, line in enumerate(lines):
        spans = [s for s in line.get("spans", []) if s.get("text")]
        for si, span in enumerate(spans):
            text = span["text"].replace("\xa0", " ")  # PDF word gaps often arrive as no-break spaces
            if si == len(spans) - 1 and li < len(lines) - 1:
                # a block's lines are one paragraph: join them with a space; a line-end hyphen (PyMuPDF reports
                # wrapped words as a soft hyphen) is dropped so the word is glued back together
                if text.endswith("\xad") or (text.endswith("-") and len(text) > 1 and text[-2].isalpha()):
                    text = text[:-1]
                else:
                    text += " "
            run = para.add_run(text)
            font_name = span.get("font", "")
            flags = span.get("flags", 0)
            lower = font_name.lower()
            run.bold = bool(flags & 16) or "bold" in lower or "black" in lower
            run.italic = bool(flags & 2) or "italic" in lower or "oblique" in lower
            run.font.size = Pt(max(1.0, round(span.get("size", 11) * 2) / 2))
            family = _family(font_name)
            if family:
                run.font.name = family
            color = span.get("color", 0)
            if color:
                run.font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)
    return y1


def _add_image(doc, block: dict, max_w: float, prev_bottom: float | None) -> float:
    data, ext = block.get("image"), (block.get("ext") or "").lower()
    x0, y0, x1, y1 = block["bbox"]
    if not data:
        return y1
    if ext not in _DOCX_IMAGES:
        from PIL import Image  # JPX, JBIG2 and friends: Word can't embed them, so re-encode

        buf = io.BytesIO()
        Image.open(io.BytesIO(data)).convert("RGB").save(buf, "PNG")
        data = buf.getvalue()
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.space_before = _gap(y0, prev_bottom)
    para.add_run().add_picture(io.BytesIO(data), width=Pt(min(max(x1 - x0, 8), max_w)))
    return y1


def _add_table(doc, rows: list[list[str | None]]) -> None:
    cols = max((len(r) for r in rows), default=0)
    if not rows or not cols:
        return
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c in range(cols):
            text = row[c] if c < len(row) and row[c] else ""
            table.cell(r, c).text = text.replace("\n", " ").strip()


def _page_items(page) -> list[tuple[float, float, str, object]]:
    """Everything worth keeping on a page as (y, x, kind, payload), in reading order."""
    items: list[tuple[float, float, str, object]] = []
    tables = []
    try:
        tables = list(page.find_tables().tables)
    except Exception:  # noqa: BLE001 - table detection is best effort; the text still comes through
        pass
    boxes = []
    for t in tables:
        rows = t.extract()
        if rows and any(any(cell for cell in row) for row in rows):
            boxes.append(pymupdf.Rect(t.bbox))
            items.append((t.bbox[1], t.bbox[0], "table", rows))

    for block in page.get_text("dict", sort=True)["blocks"]:
        if any(_inside(box, block["bbox"]) for box in boxes):
            continue
        if block["type"] == 0:
            if any(s.get("text", "").strip() for ln in block["lines"] for s in ln["spans"]):
                items.append((block["bbox"][1], block["bbox"][0], "text", block))
        elif block["type"] == 1:
            items.append((block["bbox"][1], block["bbox"][0], "image", block))
    items.sort(key=lambda i: (round(i[0]), i[1]))
    return items


def convert(src: Path, dst: Path) -> None:
    doc = Document()
    with pymupdf.open(src) as pdf:
        if pdf.needs_pass:
            raise ValueError("The PDF is password protected.")
        first = pdf[0].rect if pdf.page_count else pymupdf.Rect(0, 0, 595, 842)
        section = doc.sections[0]
        section.page_width, section.page_height = Emu(int(first.width * 12700)), Emu(int(first.height * 12700))
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Pt(_MARGIN))
        usable = max(first.width - 2 * _MARGIN, 72)

        # drop python-docx's empty default paragraph so page 1 starts at the top
        for p in list(doc.paragraphs):
            p._element.getparent().remove(p._element)

        for number, page in enumerate(pdf):
            if number:
                doc.add_page_break()
            items = _page_items(page)
            texts = [i[3]["bbox"][0] for i in items if i[2] == "text"]  # type: ignore[index]
            left_edge = min(texts) if texts else _MARGIN
            prev_bottom: float | None = None
            for _y, _x, kind, payload in items:
                if kind == "table":
                    _add_table(doc, payload)  # type: ignore[arg-type]
                    prev_bottom = None
                elif kind == "text":
                    prev_bottom = _add_text(doc, payload, page.rect.width, prev_bottom, left_edge)  # type: ignore[arg-type]
                else:
                    prev_bottom = _add_image(doc, payload, usable, prev_bottom)  # type: ignore[arg-type]
    doc.save(str(dst))
