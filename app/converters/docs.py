"""Built-in document engine: no LibreOffice, nothing to install.

Documents are read with small pure-Python libraries (mammoth, openpyxl, xlrd,
python-pptx, odfpy), turned into HTML, and printed to PDF by the system's
Edge or Chrome. If neither is installed, flowing documents fall back to
PyMuPDF's own HTML renderer.

This is a content converter, not a layout engine: text, headings, lists,
tables and images survive; exact page layout, headers/footers and embedded
fonts do not. Anything outside NATIVE_DOC_PAIRS goes to LibreOffice instead.
"""

from __future__ import annotations

import base64
import csv
import datetime as dt
import html
import io
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
from typing import Callable, Optional

from app.core.formats import NATIVE_DOC_PAIRS
from app.i18n import t

MAX_ROWS_PER_SHEET = 5000  # a 200k-row sheet would produce a PDF nobody can open

_BROWSERS = (
    r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
)

_FLOW_CSS = """
@page { size: %(size)s; margin: 18mm; }
html { -webkit-print-color-adjust: exact; print-color-adjust: exact; }
body { font-family: "Segoe UI", Calibri, Arial, sans-serif; font-size: 11pt; line-height: 1.45; color: #111; }
h1 { font-size: 20pt; margin: 0 0 8pt; } h2 { font-size: 15pt; margin: 14pt 0 6pt; }
h3 { font-size: 12.5pt; margin: 12pt 0 4pt; } p { margin: 0 0 7pt; }
img { max-width: 100%%; height: auto; }
table { border-collapse: collapse; width: 100%%; margin: 6pt 0 10pt; font-size: %(table_pt)spt; }
th, td { border: 1px solid #9aa5a0; padding: 3pt 6pt; vertical-align: top; text-align: left; overflow-wrap: anywhere; }
th { background: #eef2f0; font-weight: 600; }
pre { white-space: pre-wrap; overflow-wrap: anywhere; font-family: Consolas, "Courier New", monospace; font-size: 10pt; margin: 0; }
.sheet { break-after: page; } .sheet:last-child { break-after: auto; }
.note { color: #666; font-size: 9pt; }
"""


def find_browser() -> Optional[Path]:
    local = os.environ.get("LOCALAPPDATA")
    candidates = list(_BROWSERS)
    if local:
        candidates.append(str(Path(local) / "Google" / "Chrome" / "Application" / "chrome.exe"))
    for c in candidates:
        if Path(c).exists():
            return Path(c)
    return None


# ----------------------------------------------------------------- html -> pdf
def _print_with_browser(browser: Path, document: str, dst: Path) -> None:
    flags = subprocess.CREATE_NO_WINDOW if hasattr(subprocess, "CREATE_NO_WINDOW") else 0
    with tempfile.TemporaryDirectory(prefix="triton_html_", ignore_cleanup_errors=True) as tmp:
        tmp_dir = Path(tmp)
        page = tmp_dir / "page.html"
        out = tmp_dir / "out.pdf"
        page.write_text(document, encoding="utf-8")
        # a private profile: otherwise the command can hand off to an already-running browser and print nothing
        cmd = [
            str(browser), "--headless", "--disable-gpu", "--no-first-run", "--disable-extensions",
            f"--user-data-dir={tmp_dir / 'profile'}", "--no-pdf-header-footer", "--print-to-pdf-no-header",
            f"--print-to-pdf={out}", page.as_uri(),
        ]
        try:
            subprocess.run(cmd, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120, creationflags=flags)
        except subprocess.TimeoutExpired:
            pass  # some builds linger after writing the file; the check below is what matters
        if not out.exists() or out.stat().st_size == 0:
            raise RuntimeError(t("err_pdf_print"))
        shutil.move(str(out), str(dst))


def _print_with_story(document: str, dst: Path, landscape: bool) -> None:
    import pymupdf

    media = pymupdf.paper_rect("a4-l" if landscape else "a4")
    where = media + (40, 40, -40, -40)
    story = pymupdf.Story(html=document)
    writer = pymupdf.DocumentWriter(str(dst))
    more = 1
    while more:
        dev = writer.begin_page(media)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()


def html_to_pdf(body: str, dst: Path, *, css: Optional[str] = None, landscape: bool = False,
                table_pt: float = 10, flow: bool = True) -> None:
    """`css` replaces the default flow stylesheet (the pptx path brings its own)."""
    style = css if css is not None else _FLOW_CSS % {
        "size": "A4 landscape" if landscape else "A4", "table_pt": table_pt}
    document = f"<!doctype html><html><head><meta charset='utf-8'><style>{style}</style></head><body>{body}</body></html>"

    browser = find_browser()
    if browser:
        try:
            _print_with_browser(browser, document, dst)
            return
        except Exception:  # noqa: BLE001 - fall through to the built-in renderer
            if not flow:
                raise
    elif not flow:
        raise RuntimeError(t("err_no_browser"))
    _print_with_story(document, dst, landscape)


# ------------------------------------------------------------------- helpers
def _esc(s) -> str:
    return html.escape("" if s is None else str(s))


def read_text(path: Path) -> str:
    """Turkish files are often cp1254, not UTF-8; try UTF-8 first, then fall back."""
    raw = path.read_bytes()
    if raw[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return raw.decode("utf-16")
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1254", errors="replace")


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.datetime):
        return v.strftime("%d.%m.%Y") if v.time() == dt.time(0) else v.strftime("%d.%m.%Y %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%d.%m.%Y")
    return str(v)


Sheet = tuple[str, list[list]]


def _read_csv(path: Path) -> list[Sheet]:
    text = read_text(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [(path.stem, [row for row in csv.reader(io.StringIO(text), dialect)])]


def _read_xlsx(path: Path) -> list[Sheet]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    finally:
        wb.close()


def _read_xls(path: Path) -> list[Sheet]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    sheets: list[Sheet] = []
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype == xlrd.XL_CELL_EMPTY:
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        sheets.append((sh.name, rows))
    return sheets


_SHEET_READERS: dict[str, Callable[[Path], list[Sheet]]] = {"csv": _read_csv, "xlsx": _read_xlsx, "xls": _read_xls}


def _trim(rows: list[list]) -> list[list]:
    """Drop fully empty trailing rows/columns (spreadsheets love to report a huge used range)."""
    rows = [list(r) for r in rows]
    while rows and all(c in (None, "") for c in rows[-1]):
        rows.pop()
    width = 0
    for r in rows:
        for i in range(len(r) - 1, -1, -1):
            if r[i] not in (None, ""):
                width = max(width, i + 1)
                break
    return [r[:width] + [None] * (width - len(r[:width])) for r in rows]


def _sheets_to_html(sheets: list[Sheet], show_names: bool) -> tuple[str, int]:
    parts, widest = [], 0
    for name, rows in sheets:
        rows = _trim(rows)
        hidden = max(0, len(rows) - MAX_ROWS_PER_SHEET)
        rows = rows[:MAX_ROWS_PER_SHEET]
        widest = max(widest, max((len(r) for r in rows), default=0))
        section = ["<div class='sheet'>"]
        if show_names:
            section.append(f"<h2>{_esc(name)}</h2>")
        if rows:
            section.append("<table>")
            for i, row in enumerate(rows):
                tag = "th" if i == 0 else "td"
                section.append("<tr>" + "".join(f"<{tag}>{_esc(_cell(c))}</{tag}>" for c in row) + "</tr>")
            section.append("</table>")
        if hidden:
            section.append(f"<p class='note'>{_esc(t('docs_rows_hidden', count=hidden))}</p>")
        section.append("</div>")
        parts.append("".join(section))
    return "".join(parts), widest


# ------------------------------------------------------------------ converters
def _spreadsheet_to_pdf(src: Path, dst: Path) -> None:
    sheets = _SHEET_READERS[src.suffix.lower().lstrip(".")](src)
    body, widest = _sheets_to_html(sheets, show_names=len(sheets) > 1)
    html_to_pdf(body, dst, landscape=widest > 6, table_pt=9 if widest > 8 else 10)


def _spreadsheet_to_csv(src: Path, dst: Path) -> None:
    sheets = _SHEET_READERS[src.suffix.lower().lstrip(".")](src)
    with open(dst, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        for row in _trim(sheets[0][1]):
            w.writerow([_cell(c) for c in row])


_PLAIN_NUMBER = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")


def _to_xlsx(src: Path, dst: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in _SHEET_READERS[src.suffix.lower().lstrip(".")](src):
        ws = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "_", name)[:31] or "Sheet")
        for row in rows:
            out = []
            for c in row:
                # only unambiguous numbers become numbers: "0532..." and "00123" must stay text
                if isinstance(c, str) and _PLAIN_NUMBER.match(c):
                    c = float(c) if "." in c else int(c)
                out.append(c)
            ws.append(out)
    wb.save(dst)


def _docx_to_html(src: Path) -> str:
    import mammoth

    with open(src, "rb") as f:
        return mammoth.convert_to_html(f).value  # images are inlined as data URIs


def _docx_to_pdf(src: Path, dst: Path) -> None:
    html_to_pdf(_docx_to_html(src), dst)


def _docx_to_txt(src: Path, dst: Path) -> None:
    import mammoth

    with open(src, "rb") as f:
        dst.write_text(mammoth.extract_raw_text(f).value, encoding="utf-8")


def _txt_to_pdf(src: Path, dst: Path) -> None:
    html_to_pdf(f"<pre>{_esc(read_text(src))}</pre>", dst)


def _txt_to_docx(src: Path, dst: Path) -> None:
    from docx import Document

    doc = Document()
    for line in read_text(src).splitlines():
        doc.add_paragraph(line)
    doc.save(str(dst))


def _odt_to_pdf(src: Path, dst: Path) -> None:
    from odf.odf2xhtml import ODF2XHTML

    page = ODF2XHTML().odf2xhtml(str(src))
    # odfpy leaves images as "Pictures/x.png"; inline them from the archive so they survive
    with zipfile.ZipFile(src) as z:
        def inline(m: re.Match) -> str:
            name = m.group(1)
            try:
                data = z.read(name)
            except KeyError:
                return m.group(0)
            mime = "image/svg+xml" if name.lower().endswith(".svg") else \
                "image/" + name.rsplit(".", 1)[-1].lower().replace("jpg", "jpeg")
            return f'src="data:{mime};base64,{base64.b64encode(data).decode()}"'
        page = re.sub(r'src="(Pictures/[^"]+)"', inline, page)
    body = re.search(r"<body[^>]*>(.*)</body>", page, re.S)
    style = "".join(re.findall(r"<style[^>]*>(.*?)</style>", page, re.S))
    html_to_pdf(body.group(1) if body else page, dst, css=_FLOW_CSS % {"size": "A4", "table_pt": 10} + style)


# -------------------------------------------------------------------- pptx
def _pptx_to_pdf(src: Path, dst: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor  # noqa: F401  (ensures the colour classes are importable when frozen)
    from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

    prs = Presentation(str(src))
    px = lambda emu: round((emu or 0) / 9525, 2)  # noqa: E731 - EMU -> CSS px at 96 dpi
    width, height = px(prs.slide_width), px(prs.slide_height)

    def rgb(color) -> Optional[str]:
        try:
            if color.type == MSO_COLOR_TYPE.RGB:
                return f"#{color.rgb}"
        except Exception:  # noqa: BLE001 - theme/inherited colours aren't resolvable here
            pass
        return None

    def luminance(hex_colour: str) -> float:
        r, g, b = (int(hex_colour[i:i + 2], 16) for i in (1, 3, 5))
        return 0.299 * r + 0.587 * g + 0.114 * b

    def shapes_of(container):
        for sh in container:
            try:
                is_group = sh.shape_type == MSO_SHAPE_TYPE.GROUP
            except NotImplementedError:
                is_group = False
            if is_group:
                yield from shapes_of(sh.shapes)
            else:
                yield sh

    align_css = {PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right", PP_ALIGN.JUSTIFY: "justify"}
    anchor_css = {MSO_ANCHOR.MIDDLE: "center", MSO_ANCHOR.BOTTOM: "flex-end"}
    title_types = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
    bullet_types = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}

    slides_html = []
    for slide in prs.slides:
        bg, default_text = "#ffffff", "#111111"
        try:
            if not slide.follow_master_background and slide.background.fill.type == MSO_FILL.SOLID:
                bg = rgb(slide.background.fill.fore_color) or bg
        except Exception:  # noqa: BLE001
            pass
        if luminance(bg) < 110:
            default_text = "#f5f5f5"

        items = []
        for sh in shapes_of(slide.shapes):
            try:
                box = f"left:{px(sh.left)}px;top:{px(sh.top)}px;width:{px(sh.width)}px;height:{px(sh.height)}px;"
            except Exception:  # noqa: BLE001
                continue

            if getattr(sh, "has_table", False) and sh.has_table:
                rows = "".join(
                    "<tr>" + "".join(f"<td>{_esc(c.text)}</td>" for c in r.cells) + "</tr>" for r in sh.table.rows)
                items.append(f"<div class='shape' style='{box}'><table class='pt'>{rows}</table></div>")
                continue

            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    mime = sh.image.content_type
                    if mime.startswith("image/") and "emf" not in mime and "wmf" not in mime:
                        data = base64.b64encode(sh.image.blob).decode()
                        items.append(f"<img class='shape' style='{box}' src='data:{mime};base64,{data}'>")
                except Exception:  # noqa: BLE001
                    pass
                continue

            fill = ""
            try:
                if not sh.is_placeholder and sh.fill.type == MSO_FILL.SOLID and rgb(sh.fill.fore_color):
                    fill = f"background:{rgb(sh.fill.fore_color)};"
            except Exception:  # noqa: BLE001
                pass

            text_html = ""
            anchor = "flex-start"
            if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                ph = sh.placeholder_format.type if sh.is_placeholder else None
                anchor = anchor_css.get(sh.text_frame.vertical_anchor, "flex-start")
                paras = []
                for p in sh.text_frame.paragraphs:
                    runs = []
                    for r in p.runs:
                        size = r.font.size.pt if r.font.size else (p.font.size.pt if p.font.size else
                                                                   (36 if ph in title_types else 18))
                        css = f"font-size:{size}pt;"
                        if r.font.bold or (ph in title_types and r.font.bold is None):
                            css += "font-weight:700;"
                        if r.font.italic:
                            css += "font-style:italic;"
                        if r.font.underline:
                            css += "text-decoration:underline;"
                        if r.font.name:
                            css += f"font-family:'{r.font.name}',sans-serif;"
                        colour = rgb(r.font.color)
                        if colour:
                            css += f"color:{colour};"
                        runs.append(f"<span style=\"{css}\">{_esc(r.text)}</span>")
                    if not runs:
                        paras.append("<p>&nbsp;</p>")
                        continue
                    bullet = "• " if ph in bullet_types else ""
                    css = f"text-align:{align_css.get(p.alignment, 'left')};padding-left:{p.level * 28}px;"
                    paras.append(f"<p style=\"{css}\">{bullet}{''.join(runs)}</p>")
                text_html = "".join(paras)

            if text_html or fill:
                items.append(
                    f"<div class='shape' style='{box}{fill}display:flex;flex-direction:column;"
                    f"justify-content:{anchor};'>{text_html}</div>")

        slides_html.append(
            f"<section class='slide' style='background:{bg};color:{default_text}'>{''.join(items)}</section>")

    css = (
        f"@page {{ size: {width}px {height}px; margin: 0; }}"
        "html { -webkit-print-color-adjust: exact; print-color-adjust: exact; } body { margin: 0; }"
        f".slide {{ position: relative; width: {width}px; height: {height}px; overflow: hidden;"
        " break-after: page; font-family: 'Segoe UI', Calibri, Arial, sans-serif; }"
        ".slide:last-child { break-after: auto; }"
        ".shape { position: absolute; box-sizing: border-box; } p { margin: 0; line-height: 1.2; }"
        "table.pt { border-collapse: collapse; width: 100%; font-size: 14pt; }"
        "table.pt td { border: 1px solid #888; padding: 4px 8px; }"
    )
    html_to_pdf("".join(slides_html), dst, css=css, flow=False)


# ----------------------------------------------------------------------- entry
_CONVERTERS: dict[tuple[str, str], Callable[[Path, Path], None]] = {
    ("docx", "pdf"): _docx_to_pdf,
    ("docx", "txt"): _docx_to_txt,
    ("txt", "pdf"): _txt_to_pdf,
    ("txt", "docx"): _txt_to_docx,
    ("odt", "pdf"): _odt_to_pdf,
    ("pptx", "pdf"): _pptx_to_pdf,
    ("xlsx", "pdf"): _spreadsheet_to_pdf,
    ("xlsx", "csv"): _spreadsheet_to_csv,
    ("xls", "pdf"): _spreadsheet_to_pdf,
    ("xls", "csv"): _spreadsheet_to_csv,
    ("xls", "xlsx"): _to_xlsx,
    ("csv", "pdf"): _spreadsheet_to_pdf,
    ("csv", "xlsx"): _to_xlsx,
}

assert set(_CONVERTERS) == NATIVE_DOC_PAIRS, "formats.NATIVE_DOC_PAIRS and docs._CONVERTERS drifted apart"


def convert(src: Path, dst: Path) -> None:
    key = (src.suffix.lower().lstrip("."), dst.suffix.lower().lstrip("."))
    fn = _CONVERTERS[key]
    try:
        fn(src, dst)
    except RuntimeError:
        raise
    except Exception as exc:  # noqa: BLE001 - library errors are English and cryptic; add context
        raise RuntimeError(t("err_doc_read", detail=f"{type(exc).__name__}: {exc}")) from exc
